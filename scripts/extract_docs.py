#!/usr/bin/env python3
"""
scripts/extract_docs.py

source-docs/ 안의 pptx / docx / hwp 파일을 텍스트로 추출해
source-docs/extracted/ 에 .md 로 저장하고
source-docs/index.yaml 을 자동 업데이트합니다.

사용법:
    py scripts/extract_docs.py            # 새 파일만 처리
    py scripts/extract_docs.py --force    # 전체 재추출
    py scripts/extract_docs.py --dry-run  # 실제 저장 없이 목록만 출력

의존성:
    pip install python-pptx python-docx pyhwp pyyaml
"""

import argparse
import os
import re
import subprocess
import sys
from pathlib import Path

import yaml

# ── 경로 설정 ──────────────────────────────────────────────
ROOT = Path(__file__).parent.parent
SOURCE_DIR = ROOT / "source-docs"
EXTRACT_DIR = SOURCE_DIR / "extracted"
INDEX_FILE = SOURCE_DIR / "index.yaml"

# 건너뛸 폴더/파일 패턴
SKIP_DIRS = {"extracted", "notes"}
SKIP_FILES = {"index.yaml", "qa-log.md"}

# 블로그에 적합하지 않은 파일 키워드 (개인 정보 등)
SKIP_KEYWORDS = {"이력서", "영어공부", "정보처리기사", "전자계산기", "정처기"}


# ── day_hint 자동 감지 ──────────────────────────────────────
# 3 그룹 체계:
#   mon/wed/fri = CFD 기법 + 알고리즘 구현
#   tue/thu     = 유체역학 이론 + 유동 현상
#   sat/sun     = 논문 리뷰
# 개별 요일 키는 backwards-compat용으로 유지하되 같은 그룹 내에서 키워드 셋 공유.
_KW_CFD = [
    "FVM", "FEM", "수치해석", "CFD기법", "유한체적", "유한요소",
    "SIMPLE", "PIMPLE", "압력속도", "알고리즘", "구현", "implementation",
    "TDMA", "python", "코드", "Riemann", "HLLC", "Roe", "AUSM",
    "high-order", "Multigrid", "GMRES", "AMG", "implicit", "solver",
]
_KW_FLUID = [
    "유체역학", "경계층", "이론", "나비에", "스토크스", "navier", "stokes",
    "Reynolds", "레이놀즈", "천이", "turbulence", "난류", "와류", "vortex",
    "카르만", "cavitation", "공동", "충격파", "shock", "기화", "충돌모델",
    "LBM", "볼츠만", "열전달", "열역학", "기체역학", "상변화",
]
_KW_PAPER = [
    "논문", "paper", "review", "PINN", "arxiv", "journal",
    "paper review", "논문리뷰",
]

DAY_KEYWORDS = {
    "monday":    _KW_CFD,
    "wednesday": _KW_CFD,
    "friday":    _KW_CFD,
    "tuesday":   _KW_FLUID,
    "thursday":  _KW_FLUID,
    "saturday":  _KW_PAPER,
    "sunday":    _KW_PAPER,
}


def detect_day_hint(path: Path, text: str) -> str:
    combined = (path.stem + " " + path.parent.name + " " + text[:500]).lower()
    scores = {day: 0 for day in DAY_KEYWORDS}
    for day, keywords in DAY_KEYWORDS.items():
        for kw in keywords:
            if kw.lower() in combined:
                scores[day] += 1
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else "any"


def slugify(text: str) -> str:
    text = re.sub(r"[^\w가-힣\s-]", "", text)
    text = re.sub(r"\s+", "-", text.strip())
    return text[:60].rstrip("-")


# ── 추출 함수들 ────────────────────────────────────────────
def extract_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        slide_lines = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    slide_lines.append(t)
        if slide_lines:
            lines.append(f"## 슬라이드 {i}")
            lines.extend(slide_lines)
            lines.append("")
    return "\n".join(lines)


def extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    lines = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = para.style.name.lower()
        if "heading 1" in style:
            lines.append(f"# {t}")
        elif "heading 2" in style:
            lines.append(f"## {t}")
        elif "heading 3" in style:
            lines.append(f"### {t}")
        else:
            lines.append(t)
    return "\n".join(lines)


def extract_hwp(path: Path) -> str:
    # hwp5 Python API 사용 (preview_text.utf8)
    try:
        from hwp5 import filestructure as fs
        hwpfile = fs.Hwp5File(str(path))
        raw = hwpfile.preview_text.utf8
        return raw.decode("utf-8", errors="replace").replace("\r\n", "\n").replace("\r", "\n")
    except Exception:
        pass

    # fallback: hwp5txt subprocess
    try:
        result = subprocess.run(
            [sys.executable, "-m", "hwp5.hwp5txt", str(path)],
            capture_output=True, timeout=30
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.decode("utf-8", errors="replace")
    except (subprocess.TimeoutExpired, FileNotFoundError):
        pass

    print(f"  ⚠️  HWP 추출 실패 ({path.name})")
    return ""


EXTRACTORS = {
    ".pptx": extract_pptx,
    ".docx": extract_docx,
    ".hwp":  extract_hwp,
}


# ── index.yaml 로드/저장 ───────────────────────────────────
def load_index() -> dict:
    if INDEX_FILE.exists():
        with open(INDEX_FILE, encoding="utf-8") as f:
            return yaml.safe_load(f) or {"documents": []}
    return {"documents": []}


def save_index(data: dict):
    header = (
        "# source-docs/index.yaml\n"
        "# 자동 포스트 생성에 사용할 소스 문서 목록\n"
        "# day_hint: monday | tuesday | wednesday | thursday | friday | saturday | sunday | any\n"
        "# last_used: 마지막 사용 날짜 (YYYY-MM-DD), 문서 순환에 사용\n\n"
    )
    footer = (
        "\n# 새 문서 추가 방법:\n"
        "# 1. source-docs/notes/ 에 .md 파일 저장\n"
        "# 2. 아래 형식으로 항목 추가 (last_used: null 이면 가장 먼저 사용)\n"
        "#\n"
        "# - id: 고유-id\n"
        "#   title: \"문서 제목\"\n"
        "#   path: notes/파일명.md\n"
        "#   type: markdown\n"
        "#   day_hint: monday\n"
        "#   tags: [태그1, 태그2]\n"
        "#   last_used: null\n"
    )
    body = yaml.dump(
        data, allow_unicode=True, default_flow_style=False,
        sort_keys=False, indent=2
    )
    with open(INDEX_FILE, "w", encoding="utf-8") as f:
        f.write(header + body + footer)


# ── 메인 처리 ──────────────────────────────────────────────
def should_skip(path: Path) -> bool:
    name = path.name
    if any(kw in name for kw in SKIP_KEYWORDS):
        return True
    # 임시/충돌 파일
    if "_Conflict" in name or name.startswith("~$"):
        return True
    return False


def process_file(path: Path, existing_ids: set, dry_run: bool, force: bool) -> dict | None:
    suffix = path.suffix.lower()
    if suffix not in EXTRACTORS:
        return None
    if should_skip(path):
        print(f"  ⏭️  건너뜀 (키워드 필터): {path.name}")
        return None

    # extracted 경로 계산 (source-docs 기준 상대 경로 유지)
    rel = path.relative_to(SOURCE_DIR)
    out_path = EXTRACT_DIR / rel.with_suffix(".md")

    # 이미 추출됐고 force 아니면 스킵
    doc_id = slugify(str(rel.with_suffix("")).replace("\\", "-").replace("/", "-").replace(" ", "-"))
    if doc_id in existing_ids and not force:
        return None

    print(f"  📄 추출 중: {rel}")

    if dry_run:
        return {"id": doc_id, "path": str(rel)}

    # 추출
    extractor = EXTRACTORS[suffix]
    try:
        text = extractor(path)
    except Exception as e:
        print(f"  ❌ 오류: {e}")
        return None

    if not text or len(text.strip()) < 50:
        print(f"  ⚠️  내용 없음 또는 너무 짧음, 건너뜀")
        return None

    # .md 저장
    out_path.parent.mkdir(parents=True, exist_ok=True)
    title = path.stem
    md_content = f"# {title}\n\n> 원본: `{rel}`\n\n{text.strip()}\n"
    out_path.write_text(md_content, encoding="utf-8")

    # index.yaml 항목 생성
    day_hint = detect_day_hint(path, text)
    tags = [path.parent.name.strip("[] ")] if path.parent != SOURCE_DIR else []
    tags = [t for t in tags if t and t not in {"공부", "notes", "extracted"}]

    return {
        "id": doc_id,
        "title": title,
        "path": f"extracted/{rel.with_suffix('.md')}".replace("\\", "/"),
        "type": "markdown",
        "day_hint": day_hint,
        "tags": tags or ["CFD"],
        "last_used": None,
    }


def main():
    parser = argparse.ArgumentParser(description="source-docs 문서 추출기")
    parser.add_argument("--force", action="store_true", help="이미 추출된 파일도 재추출")
    parser.add_argument("--dry-run", action="store_true", help="실제 저장 없이 목록만 출력")
    args = parser.parse_args()

    EXTRACT_DIR.mkdir(exist_ok=True)

    index_data = load_index()
    existing_ids = {doc["id"] for doc in index_data.get("documents", [])}

    new_entries = []
    stats = {"pptx": 0, "docx": 0, "hwp": 0, "skipped": 0, "error": 0}

    print(f"\n🔍 source-docs/ 스캔 중...\n")

    for path in sorted(SOURCE_DIR.rglob("*")):
        if not path.is_file():
            continue
        # 건너뛸 폴더
        if any(part in SKIP_DIRS for part in path.parts):
            continue
        suffix = path.suffix.lower()
        if suffix not in EXTRACTORS:
            continue

        result = process_file(path, existing_ids, args.dry_run, args.force)
        if result:
            new_entries.append(result)
            stats[suffix.lstrip(".")] = stats.get(suffix.lstrip("."), 0) + 1
        else:
            stats["skipped"] += 1

    if not args.dry_run and new_entries:
        index_data["documents"].extend(new_entries)
        save_index(index_data)
        print(f"\n✅ index.yaml 업데이트 완료 (+{len(new_entries)}개 항목)")

    print(f"\n📊 결과:")
    print(f"   pptx: {stats['pptx']}개  docx: {stats['docx']}개  hwp: {stats['hwp']}개  건너뜀: {stats['skipped']}개")

    if args.dry_run:
        print(f"\n[dry-run] 실제 저장하려면 --dry-run 없이 실행하세요.")
    else:
        print(f"\n추출된 파일 위치: source-docs/extracted/")
        print(f"다음 단계: index.yaml에서 day_hint 확인 후 조정하세요.")


if __name__ == "__main__":
    main()
